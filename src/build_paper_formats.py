"""Build the single paper's TeX and Word representations from paper.md."""
from pathlib import Path
import subprocess
import re
import argparse
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source-root', type=Path, default=Path(__file__).resolve().parents[1])
parser.add_argument('--output-root', type=Path, required=True,
                    help='New directory for rebuilt formats; source manuscript is read-only')
parser.add_argument('--work-root', type=Path, required=True)
parser.add_argument('--pandoc', type=Path, required=True)
parser.add_argument('--xelatex', type=Path)
args = parser.parse_args()
SOURCE_REPO = args.source_root.resolve()
REPO = args.output_root.resolve()
WORK = args.work_root.resolve()
PANDOC = args.pandoc.resolve()
if args.xelatex:
    args.xelatex = args.xelatex.resolve()
for output in (REPO, WORK):
    if output.exists():
        raise FileExistsError('Use a new build directory: ' + str(output))
    for protected in (SOURCE_REPO, Path(__file__).resolve().parents[1]):
        if output == protected or output.is_relative_to(protected) or protected.is_relative_to(output):
            raise ValueError('Build directories must be outside the source and executing packages')
if REPO == WORK or REPO.is_relative_to(WORK) or WORK.is_relative_to(REPO):
    raise ValueError('Output and build-work directories must be separate')
WORK.mkdir(parents=True)
(REPO / 'manuscript').mkdir(parents=True)
shutil.copy2(SOURCE_REPO / 'manuscript/paper.md', REPO / 'manuscript/paper.md')
shutil.copytree(SOURCE_REPO / 'manuscript/figures', REPO / 'manuscript/figures')

default_reference = subprocess.run([str(PANDOC), '--print-default-data-file=reference.docx'], check=True, capture_output=True).stdout
(WORK / 'pandoc_reference.docx').write_bytes(default_reference)
reference = Document(WORK / 'pandoc_reference.docx')
sec = reference.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.8)
sec.left_margin = sec.right_margin = Inches(.85)
sec.header_distance = sec.footer_distance = Inches(.35)
for name in ['Normal', 'Body Text', 'First Paragraph', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    if name not in reference.styles:
        continue
    style = reference.styles[name]
    style.font.name = 'Times New Roman'; style.font.color.rgb = RGBColor(0, 0, 0)
    for fonts in style.element.xpath('./w:rPr/w:rFonts'):
        for key in list(fonts.attrib):
            if key.endswith('Theme') or key.endswith('theme'):
                del fonts.attrib[key]
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.04
    if name.startswith('Heading'):
        style.font.bold = True
        style.font.size = Pt(14 if name == 'Heading 1' else 12)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.keep_with_next = True
    if name == 'Title':
        style.font.size = Pt(20); style.font.bold = True
        style.paragraph_format.space_after = Pt(10)
        style.paragraph_format.keep_with_next = True
    if name == 'Caption':
        style.font.size = Pt(9.5); style.paragraph_format.space_after = Pt(8)
footer = sec.footer.paragraphs[0]
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
footer.add_run('UCP Technology LLC\t').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
reference.core_properties.author = 'UCP Technology LLC'
reference.save(WORK / 'reference.docx')

source = (REPO / 'manuscript/paper.md').read_text(encoding='utf-8')
assert '{{' not in source
source = re.sub(r'(?<![<(])(https?://[^\s>]+)',
                lambda m: '<' + m[1].rstrip('.') + '>' + ('.' if m[1].endswith('.') else ''), source)
# Pandoc table captions preserve numbered titles and their associated tables.
source = re.sub(r'(?m)^(Table [A-Z]?\d+\. [^\n]+)\n\n((?:\|[^\n]*\n)+)',
                lambda m: m[2] + '\n: ' + m[1] + '\n', source)
source_path = WORK / 'paper_conversion.md'
source_path.write_text(source, encoding='utf-8')
preamble = r'''
\usepackage{caption}
\usepackage{float}
\usepackage{needspace}
\usepackage{xurl}
\captionsetup{labelformat=empty,font=small}
\usepackage{fancyhdr}
\pagestyle{fancy}
\fancyhf{}
\fancyfoot[L]{\small UCP Technology LLC}
\fancyfoot[R]{\small\thepage}
\renewcommand{\headrulewidth}{0pt}
\renewcommand{\footrulewidth}{0.3pt}
\fancypagestyle{plain}{\fancyhf{}\fancyfoot[L]{\small UCP Technology LLC}\fancyfoot[R]{\small\thepage}\renewcommand{\headrulewidth}{0pt}\renewcommand{\footrulewidth}{0.3pt}}
\setlength{\emergencystretch}{3em}
\setlength{\parindent}{0pt}
\setlength{\parskip}{5pt}
\widowpenalty=10000
\clubpenalty=10000
\setcounter{topnumber}{3}
\setcounter{bottomnumber}{2}
\renewcommand{\topfraction}{0.9}
\renewcommand{\textfraction}{0.08}
\renewcommand{\floatpagefraction}{0.8}
'''
(WORK / 'preamble.tex').write_text(preamble, encoding='utf-8')
common = [str(PANDOC), str(source_path), '--from=markdown+tex_math_dollars+implicit_figures', '--standalone',
          '--resource-path=' + str(REPO / 'manuscript')]
subprocess.run(common + ['--to=latex', '-o', str(REPO / 'manuscript/paper.tex'),
    '--pdf-engine=xelatex', '-V', 'documentclass=article', '-V', 'fontsize=11pt', '-V', 'papersize=letter',
    '-V', 'geometry:top=0.8in,bottom=0.8in,left=0.85in,right=0.85in', '-V', 'mainfont=Times New Roman',
    '-V', 'sansfont=Arial', '-V', 'monofont=Courier New', '-V', 'colorlinks=true', '-V', 'urlcolor=black',
    '-V', 'linkcolor=black', '-V', 'linestretch=1.04', '--include-in-header=' + str(WORK / 'preamble.tex')],
    check=True, cwd=REPO / 'manuscript')
tex_path = REPO / 'manuscript/paper.tex'
tex = tex_path.read_text(encoding='utf-8')
# Keep figures with the evidence they illustrate, and each compact table intact.
tex = tex.replace(r'\begin{figure}', r'\begin{figure}[H]')
table_heights = iter([2.7, 2.1, 3.3, 4.1, 2.2, 4.4, 2.0, 2.1])
tex = re.sub(r'\\begin\{longtable\}', lambda m: r'\Needspace{' + str(next(table_heights)) + 'in}\n' + m[0], tex)
tex = re.sub(r'(\\\[\s*\\theta_\{min\}.*?\\\]\s*\\\[.*?\\theta_\{max\}.*?\\\])',
             lambda m: r'\begin{samepage}' + '\n' + m[0] + '\n' + r'\end{samepage}', tex, flags=re.S)
tex = re.sub(r'(?<![\w{])-(\d+\.\d+)', lambda m: r'\mbox{-' + m[1] + '}', tex)
tex_path.write_text(tex, encoding='utf-8')
subprocess.run(common + ['--to=docx', '-o', str(REPO / 'manuscript/paper.docx'),
    '--reference-doc=' + str(WORK / 'reference.docx')], check=True, cwd=REPO / 'manuscript')
doc = Document(REPO / 'manuscript/paper.docx')
doc.core_properties.author = 'UCP Technology LLC'
doc.core_properties.last_modified_by = 'UCP Technology LLC'
for p in doc.paragraphs:
    if p.text.startswith('Table '):
        p.paragraph_format.keep_with_next = True
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
    for r in p.runs:
        if p.style.name in {'Title', 'Subtitle'} or p.style.name.startswith('Heading'):
            r.font.color.rgb = RGBColor(0, 0, 0)
for p, following in zip(doc.paragraphs, doc.paragraphs[1:]):
    if p._p.xpath('./m:oMathPara') and following._p.xpath('./m:oMathPara'):
        p.paragraph_format.keep_with_next = True
for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column in table.columns:
        column.width = Inches(6.8 / len(table.columns))
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        item = OxmlElement('w:' + edge); item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), '4'); item.set(qn('w:color'), 'D9D9D9'); borders.append(item)
    table._tbl.tblPr.append(borders)
    for i, row in enumerate(table.rows):
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); pr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); pr.append(repeat)
        for cell in row.cells:
            cell.width = Inches(6.8 / len(table.columns))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EDF2' if i == 0 else ('F7F7F7' if i % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', '70'), ('bottom', '70'), ('left', '80'), ('right', '80')]:
                item = OxmlElement('w:' + edge); item.set(qn('w:w'), value); item.set(qn('w:type'), 'dxa'); margins.append(item)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = 'Times New Roman'; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 0, 0)
                    if i == 0: r.bold = True
doc.save(REPO / 'manuscript/paper.docx')
print('Created paper.tex and paper.docx from one authoritative paper.md')

if args.xelatex:
    for _ in range(2):
        with (WORK/'latex_build.log').open('w',encoding='utf-8') as stream:
            subprocess.run([str(args.xelatex),'--interaction=nonstopmode','--halt-on-error',
                '-output-directory='+str(WORK),'paper.tex'],cwd=REPO/'manuscript',
                stdout=stream,stderr=subprocess.STDOUT,check=True)
    shutil.copy2(WORK/'paper.pdf',REPO/'manuscript/paper.pdf')
    print('Created designated final paper.pdf')
